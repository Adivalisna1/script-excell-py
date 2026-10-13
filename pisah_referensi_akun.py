import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import argparse
import os
import sys
import glob

# Mapping label & nama file per Tipe Akun
TIPE_MAP = {
    "A": {
        "nama": "Aset",
        "filename": "Referensi_Akun_Tipe_A_Aset.xlsx",
        "deskripsi": "Aset (Kas, Piutang, Persediaan, Aset Tetap, dll.)"
    },
    "L": {
        "nama": "Liabilitas",
        "filename": "Referensi_Akun_Tipe_L_Liabilitas.xlsx",
        "deskripsi": "Kewajiban / Utang Jangka Pendek & Panjang"
    },
    "O": {
        "nama": "Ekuitas",
        "filename": "Referensi_Akun_Tipe_O_Ekuitas.xlsx",
        "deskripsi": "Ekuitas / Modal Bersih Pemerintah"
    },
    "R": {
        "nama": "Pendapatan",
        "filename": "Referensi_Akun_Tipe_R_Pendapatan.xlsx",
        "deskripsi": "Pendapatan Pajak, PNBP, Hibah & Transfer"
    },
    "E": {
        "nama": "Beban_dan_Belanja",
        "filename": "Referensi_Akun_Tipe_E_Beban_dan_Belanja.xlsx",
        "deskripsi": "Belanja Pegawai, Barang, Modal, Bansos & Beban"
    }
}

def pisah_referensi_akun(file_input, folder_output="output_per_tipe_akun"):
    if not os.path.exists(file_input):
        print(f"[!] Error: File '{file_input}' tidak ditemukan.")
        sys.exit(1)

    print("=" * 70)
    print("   PROGRAM PEMISAHAN REFERENSI AKUN PER TIPE AKUN")
    print("=" * 70)
    print(f"[*] File Sumber : {os.path.basename(file_input)}")
    print(f"[*] Folder Hasil: {folder_output}\n")

    os.makedirs(folder_output, exist_ok=True)

    print("[*] Membaca data file Excel sumber...")
    wb_source = openpyxl.load_workbook(file_input, data_only=True)
    sheet_source = wb_source.active

    all_rows = list(sheet_source.iter_rows(values_only=True))
    if len(all_rows) < 5:
        print("[!] Error: Format file tidak sesuai, jumlah baris terlalu sedikit.")
        sys.exit(1)

    # Baris 5 (index 4) adalah header
    header_titles = ["NO", "Kode Akun", "Nama Akun", "Tipe Akun"]

    grouped_data = {k: [] for k in TIPE_MAP.keys()}
    grouped_data["LAINNYA"] = []

    total_valid_rows = 0
    for r in all_rows[5:]:
        if not r or all(cell is None for cell in r):
            continue
        kode_akun = str(r[1]).strip() if r[1] is not None else ""
        nama_akun = str(r[2]).strip() if r[2] is not None else ""
        tipe_akun = str(r[3]).strip().upper() if r[3] is not None else ""

        if not kode_akun and not nama_akun:
            continue

        item = (kode_akun, nama_akun, tipe_akun)
        if tipe_akun in grouped_data:
            grouped_data[tipe_akun].append(item)
        else:
            grouped_data["LAINNYA"].append(item)
        total_valid_rows += 1

    print(f"[*] Total baris akun berhasil dibaca: {total_valid_rows:,} akun.\n")

    # Styling Excel profesional
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="173B75", end_color="173B75", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    
    thin_border = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    summary_results = []

    for kode_tipe, info in TIPE_MAP.items():
        data_list = grouped_data[kode_tipe]
        target_path = os.path.join(folder_output, info["filename"])

        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = f"Tipe {kode_tipe} - {info['nama']}"

        # Header Title Info
        ws_out.append([f"Referensi Akun - Tipe {kode_tipe} ({info['nama']})"])
        ws_out.cell(row=1, column=1).font = Font(name="Calibri", size=13, bold=True, color="173B75")
        ws_out.append([f"Kategori: {info['deskripsi']} | Total: {len(data_list):,} Akun"])
        ws_out.cell(row=2, column=1).font = Font(name="Calibri", size=10, italic=True, color="555555")
        ws_out.append([])

        # Table Header
        ws_out.append(header_titles)
        header_row_num = 4
        for col_idx in range(1, 5):
            cell = ws_out.cell(row=header_row_num, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = cell_border

        # Data Rows
        for seq, (kode, nama, tipe) in enumerate(data_list, start=1):
            row_num = header_row_num + seq
            ws_out.append([seq, kode, nama, tipe])

            c_no = ws_out.cell(row=row_num, column=1)
            c_no.alignment = center_align
            c_no.font = regular_font
            c_no.border = cell_border

            c_kode = ws_out.cell(row=row_num, column=2)
            c_kode.alignment = center_align
            c_kode.font = regular_font
            c_kode.border = cell_border
            c_kode.number_format = '@'

            c_nama = ws_out.cell(row=row_num, column=3)
            c_nama.alignment = left_align
            c_nama.font = regular_font
            c_nama.border = cell_border

            c_tipe = ws_out.cell(row=row_num, column=4)
            c_tipe.alignment = center_align
            c_tipe.font = regular_font
            c_tipe.border = cell_border

        ws_out.column_dimensions['A'].width = 8
        ws_out.column_dimensions['B'].width = 16
        ws_out.column_dimensions['C'].width = 55
        ws_out.column_dimensions['D'].width = 14

        wb_out.save(target_path)
        file_size_kb = os.path.getsize(target_path) / 1024
        summary_results.append({
            "tipe": kode_tipe,
            "nama": info["nama"],
            "filename": info["filename"],
            "jumlah": len(data_list),
            "size": f"{file_size_kb:.1f} KB"
        })
        print(f"[OK] Berhasil membuat: {info['filename']} ({len(data_list):,} akun, {file_size_kb:.1f} KB)")

    if grouped_data["LAINNYA"]:
        other_path = os.path.join(folder_output, "Referensi_Akun_Lainnya.xlsx")
        wb_out = openpyxl.Workbook()
        ws_out = wb_out.active
        ws_out.title = "Lainnya"
        ws_out.append(header_titles)
        for seq, (kode, nama, tipe) in enumerate(grouped_data["LAINNYA"], start=1):
            ws_out.append([seq, kode, nama, tipe])
        wb_out.save(other_path)
        print(f"[OK] Berhasil membuat: Referensi_Akun_Lainnya.xlsx ({len(grouped_data['LAINNYA']):,} akun)")

    print("\n" + "=" * 70)
    print("                      RINGKASAN PEMISAHAN FILE")
    print("=" * 70)
    print(f"{'Tipe':<6} | {'Kategori':<20} | {'Jumlah Akun':<12} | {'Nama File Output'}")
    print("-" * 70)
    for res in summary_results:
        print(f"[{res['tipe']}]    | {res['nama']:<20} | {res['jumlah']:>6,} akun  | {res['filename']}")
    print("-" * 70)
    print(f"Total Akun Terproses: {total_valid_rows:,} akun.")
    print(f"Semua file tersimpan di folder:\n-> {folder_output}")
    print("=" * 70)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Pilah Referensi Akun Excel menjadi 5 file terpisah per Tipe Akun (A, L, O, R, E).")
    parser.add_argument("file_input", nargs="?", default=None, help="Path ke file Referensi Akun Excel mentahan.")
    parser.add_argument("-o", "--output", default="output_per_tipe_akun", help="Nama folder tempat menyimpan file hasil (default: output_per_tipe_akun).")

    args = parser.parse_args()

    input_file = args.file_input
    if not input_file:
        candidates = glob.glob("*Referensi*Akun*.xlsx") + glob.glob("../*Referensi*Akun*.xlsx") + glob.glob("../data_excel/*Referensi*Akun*.xlsx")
        if candidates:
            input_file = candidates[0]
            print(f"[*] Menggunakan file terdeteksi otomatis: {input_file}")
        else:
            print("[!] Harap cantumkan file Excel: python pisah_referensi_akun.py <nama_file.xlsx>")
            sys.exit(1)

    pisah_referensi_akun(input_file, args.output)
