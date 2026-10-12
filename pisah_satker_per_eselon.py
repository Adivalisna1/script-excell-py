import os
import sys
import argparse
import re
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def bersihkan_nama_file(nama):
    # Hilangkan karakter ilegal untuk nama file di Windows
    nama_bersih = re.sub(r'[\\/*?:"<>|]', '', str(nama)).strip()
    # Ganti spasi ganda dengan spasi tunggal
    nama_bersih = re.sub(r'\s+', ' ', nama_bersih)
    return nama_bersih

def format_excel_worksheet(ws):
    # Styling header
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_align = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )
    
    stripe_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    ws.row_dimensions[1].height = 28

    # Format header baris 1
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align
        cell.border = thin_border

    # Format baris data
    for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
        ws.row_dimensions[row_idx].height = 20
        fill_to_use = stripe_fill if row_idx % 2 == 0 else white_fill
        for cell in row:
            cell.border = thin_border
            cell.fill = fill_to_use
            # Align center untuk kode satker / kode eselon / kategori
            if cell.column_letter in ['A', 'C', 'G', 'I', 'J', 'K', 'L']:
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column width
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def pisahkan_satker_per_eselon(file_input="script excell/kkpk satker.xlsx", folder_output="script excell/output_per_eselon1"):
    if not os.path.exists(file_input):
        # Coba fallback jika dijalankan dari dalam folder 'script excell'
        if os.path.exists("kkpk satker.xlsx"):
            file_input = "kkpk satker.xlsx"
            if folder_output == "script excell/output_per_eselon1":
                folder_output = "output_per_eselon1"
        else:
            print(f"[!] Error: File input '{file_input}' tidak ditemukan.")
            sys.exit(1)

    print(f"[*] Membaca file master: '{file_input}'...")
    try:
        excel_file = pd.ExcelFile(file_input)
        target_sheet = 'klpk satker' if 'klpk satker' in excel_file.sheet_names else excel_file.sheet_names[0]
        df = pd.read_excel(file_input, sheet_name=target_sheet)
        print(f"[*] Menggunakan sheet '{target_sheet}' (Total {len(df)} baris satker).")
    except Exception as e:
        print(f"[!] Gagal membaca Excel: {e}")
        sys.exit(1)

    # Identifikasi kolom Eselon 1
    kolom_eselon = None
    for target in ['NAMA ESELON 1', 'eselon1', 'ESELON 1', 'nama_eselon_1']:
        for col in df.columns:
            if str(col).strip().lower() == target.lower():
                kolom_eselon = col
                break
        if kolom_eselon:
            break

    if not kolom_eselon:
        print("[!] Error: Kolom 'NAMA ESELON 1' atau 'eselon1' tidak ditemukan dalam data.")
        sys.exit(1)

    # Bersihkan kode satker agar tidak ada akhiran desimal .0
    if 'KDSATKER' in df.columns:
        df['KDSATKER'] = df['KDSATKER'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
    if 'ESELON 1' in df.columns:
        df['ESELON 1'] = df['ESELON 1'].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()

    os.makedirs(folder_output, exist_ok=True)
    print(f"[*] Folder output: '{folder_output}'\n")

    # Kelompokkan data per Eselon 1
    grouped = df.groupby(kolom_eselon, dropna=False)
    daftar_hasil = []

    for nama_eselon_raw, group_df in grouped:
        nama_eselon = "Eselon 1 Lainnya" if pd.isna(nama_eselon_raw) or not str(nama_eselon_raw).strip() else str(nama_eselon_raw).strip()
        nama_bersih = bersihkan_nama_file(nama_eselon)
        
        # Ambil kode eselon jika ada
        kode_eselon = ""
        if 'ESELON 1' in group_df.columns and not group_df['ESELON 1'].empty:
            kode_first = str(group_df['ESELON 1'].iloc[0]).strip()
            if kode_first and kode_first != "nan":
                kode_eselon = f"[{kode_first}] "

        nama_file_excel = f"{kode_eselon}{nama_bersih}.xlsx"
        path_output = os.path.join(folder_output, nama_file_excel)

        # Simpan ke Excel menggunakan openpyxl engine
        with pd.ExcelWriter(path_output, engine='openpyxl') as writer:
            group_df.to_excel(writer, sheet_name="Daftar Satker", index=False)

        # Berikan styling visual profesional
        wb = openpyxl.load_workbook(path_output)
        ws = wb["Daftar Satker"]
        format_excel_worksheet(ws)
        wb.save(path_output)

        jumlah_satker = len(group_df)
        daftar_hasil.append((nama_eselon, nama_file_excel, jumlah_satker))
        print(f"  [OK] Dibuat: '{nama_file_excel}' -> {jumlah_satker} Satker")

    print("\n" + "="*70)
    print("PEMILAHAN SATKER BERDASARKAN ESELON 1 SELESAI!")
    print("="*70)
    print(f"{'No':<4} | {'Unit Eselon 1':<45} | {'Jumlah Satker':<15}")
    print("-" * 70)
    total_satker = 0
    for idx, (eselon, fname, count) in enumerate(daftar_hasil, start=1):
        print(f"{idx:<4} | {eselon[:45]:<45} | {count} Satker")
        total_satker += count
    print("-" * 70)
    print(f"Total: {len(daftar_hasil)} File Excel Berhasil Dihasilkan ({total_satker} Satker)")
    print(f"Lokasi File: {os.path.abspath(folder_output)}")
    print("="*70)

if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Script untuk memilah master satker per Eselon 1 ke file Excel terpisah.")
    parser.add_argument("-i", "--input", default="script excell/kkpk satker.xlsx", help="File Excel sumber master satker")
    parser.add_argument("-o", "--output", default="script excell/output_per_eselon1", help="Folder tujuan untuk menyimpan file hasil")
    args = parser.parse_args()

    pisahkan_satker_per_eselon(args.input, args.output)
